from pathlib import Path

from pptx import Presentation


def add_bullet_slide(prs: Presentation, title: str, bullets: list[str], notes: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title + Content
    slide.shapes.title.text = title
    body = slide.shapes.placeholders[1].text_frame
    body.clear()
    for i, bullet in enumerate(bullets):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = bullet
        p.level = 0
    notes_tf = slide.notes_slide.notes_text_frame
    notes_tf.clear()
    notes_tf.text = " ".join(notes)


def main() -> None:
    prs = Presentation()
    prs.core_properties.title = "CS455 Traffic Simulation Project"
    prs.core_properties.subject = "Data Structures and Algorithms Project Presentation"
    prs.core_properties.author = "CS455 Project Team"

    slides: list[tuple[str, list[str], list[str]]] = [
        (
            "Project Overview & Problem",
            [
                "Python/Pygame traffic simulation in car_testing_env.py.",
                "Core problem: route multiple cars on a shared grid without collisions or deadlock.",
                "Environment includes a mandatory 4-way intersection and a parking-lot chokepoint.",
                "Goal condition: all tracked cars reach correct destinations.",
                "Focus: demonstrate data structures + algorithms in an interactive system.",
            ],
            [
                "This project models traffic as a graph problem on a 2D grid.",
                "The hard part is coordinating several agents that compete for the same road tiles.",
                "We used BFS, queues, sets, and dictionaries to solve routing and coordination issues.",
                "The final simulation is both playable and analyzable.",
            ],
        ),
        (
            "Data Structures Used",
            [
                "List[List[int]] grid stores ROAD/WALL map for O(1) cell access.",
                "List[Car] stores all cars and their state (pos, target, policy, move_tick).",
                "deque used for BFS frontier and rolling simulation logs.",
                "set tracks visited nodes, occupied tiles, blocked tiles, and waiting directions.",
                "dict parent map reconstructs shortest paths after BFS reaches goal.",
            ],
            [
                "Each structure was chosen for runtime behavior.",
                "deque gives fast popleft for BFS.",
                "sets make occupancy and visited checks efficient.",
                "The parent dictionary turns reachability into a concrete route.",
            ],
        ),
        (
            "Algorithms Implemented",
            [
                "BFS traversal computes shortest path on an unweighted 4-neighbor grid.",
                "Path reconstruction walks parent pointers from goal back to start.",
                "Right-of-way scheduler controls entry into the center intersection tile.",
                "Support-car policy intentionally moves away from congestion center.",
                "Optimized mode adds periodic random stress; god mode stays deterministic BFS.",
            ],
            [
                "BFS guarantees shortest path length in this map because every move has equal cost.",
                "Parent backtracking gives the next executable move for each agent.",
                "We added intersection scheduling because shortest paths alone cannot resolve simultaneous conflicts.",
            ],
        ),
        (
            "System Functionality & Controls",
            [
                "Turn-based updates: AI moves after player move or stall.",
                "Controls: WASD/arrows move, Space stall, G toggle mode, R regenerate map.",
                "God mode: deterministic BFS + player auto-step when pressing Space.",
                "Optimized mode: BFS with periodic randomness for stress testing.",
                "End-state incident report summarizes avoided conflicts and run status.",
            ],
            [
                "Turn-based steps make behavior reproducible for demos and debugging.",
                "Each input triggers one simulation tick and then AI decisions.",
                "Mode toggle lets us compare deterministic behavior versus stress behavior.",
            ],
        ),
        (
            "Design Decisions",
            [
                "Center 4-way lock forces contention and tests fairness logic.",
                "Parking lot is routed through a chokepoint for realistic bottlenecks.",
                "Approach/gate zones protect player access to destination.",
                "Mixed AI roles: BFS goal cars + one support congestion-clearing car.",
                "Fullscreen HUD exposes mode, car counts, logs, and completion state.",
            ],
            [
                "We intentionally designed constraints that create algorithmic pressure.",
                "The support car adds a nontraditional role focused on flow, not destination racing.",
                "These choices increase originality and produce richer demo behavior.",
            ],
        ),
        (
            "Challenges & Solutions",
            [
                "Challenge: simultaneous entry attempts at the 4-way center.",
                "Solution: rotating direction priority + waiting-set tracking.",
                "Challenge: player blocked near final parking corridor.",
                "Solution: gate-zone restrictions and west-retreat behavior for AI.",
                "Challenge: stale paths in dynamic traffic.",
            ],
            [
                "The biggest issue was multi-agent conflict on the same intersection tile.",
                "Explicit right-of-way state fixed fairness and prevented deadlocks.",
                "Replanning with BFS each turn kept routing robust as occupancy changed.",
            ],
        ),
        (
            "Complexity, Testing, and Results",
            [
                "BFS per car per turn: O(V + E), effectively O(R*C) on this grid.",
                "Space per BFS run: O(R*C) for visited/parent tracking.",
                "Multi-car step cost scales with number of active BFS cars.",
                "Tests: mode toggles, map regeneration, stall behavior, intersection contention, completion.",
                "Result: cars reach destinations reliably; conflicts are logged instead of causing invalid moves.",
            ],
            [
                "This project prioritizes correctness and clarity over premature optimization.",
                "On this grid, recomputing BFS each turn is practical and stable.",
                "Testing both deterministic and stress modes gave good coverage.",
            ],
        ),
        (
            "Originality, Takeaways, and Next Steps",
            [
                "Original blend: shortest-path routing + traffic-style right-of-way scheduling.",
                "Interactive player-in-the-loop stress tests algorithm behavior live.",
                "Support-car policy adds emergent congestion-management behavior.",
                "Incident-report UX turns each run into reviewable evidence.",
                "Future work: weighted edges, A* comparison, multi-intersection scaling, replay export.",
            ],
            [
                "This goes beyond a basic BFS visualizer by combining routing and coordination.",
                "Player interaction quickly exposes edge cases during demos.",
                "Key takeaway: simple data structures can produce robust multi-agent behavior when composed well.",
            ],
        ),
    ]

    for title, bullets, notes in slides:
        add_bullet_slide(prs, title, bullets, notes)

    # Demo plan slide
    add_bullet_slide(
        prs,
        "60-Second Demo Plan",
        [
            "0:00-0:05 Run: py car_testing_env.py and introduce the simulator.",
            "0:05-0:15 Press D several times; explain one keypress = one turn.",
            "0:15-0:28 Approach center and press Space to show stall behavior.",
            "0:28-0:33 Press G to switch to god mode.",
            "0:33-1:00 Press Space repeatedly; show auto-step BFS and completion.",
        ],
        [
            "This sequence is designed to reliably show routing, scheduling, and mode behavior within one minute.",
            "If completion does not occur in one minute, press R and rerun from the start.",
        ],
    )

    # Q&A slide
    add_bullet_slide(
        prs,
        "Q&A Backup",
        [
            "Why BFS instead of A*? Unweighted grid so BFS is shortest and simpler.",
            "How avoid deadlock at 4-way? Rotating priority plus waiting-direction tracking.",
            "What is complexity per turn? Roughly O(k*R*C) for k BFS cars.",
            "How prevent collisions? Occupancy checks block invalid moves and log incidents.",
            "Why re-run BFS each turn? Dynamic occupancy changes shortest paths.",
            "What is most original? Combining pathfinding with traffic-style right-of-way control.",
        ],
        [
            "These are the most common questions from project demos.",
            "Each answer maps directly to implementation details in the code.",
        ],
    )

    out = Path("Project_Presentation.pptx")
    prs.save(out)
    print(f"Created {out.resolve()}")


if __name__ == "__main__":
    main()
